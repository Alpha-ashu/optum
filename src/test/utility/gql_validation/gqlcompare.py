import os
import sys
import json
import logging
import re
from collections import defaultdict
import pandas as pd
from typing import Any, List, Dict
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import FormulaRule

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

var = sys.argv[1] if len(sys.argv) > 1 else "ClaimTransactionIdentifier"
logging.info(f"Consumer argument received: {var}")

dir_path = os.getcwd().split("src")[0]

alex_dir = os.path.join(dir_path, 'target', 'All_Responses', 'Alex_Responses')
ppkg_dir = os.path.join(dir_path, 'target', 'All_Responses', 'GQL_Responses')
output_dir = os.path.join(dir_path, 'target', 'All_Responses', 'Filtered_Responses')
os.makedirs(output_dir, exist_ok=True)
output_file = os.path.join(output_dir, f"{var}_validated_report.xlsx")

if os.path.exists(output_file):
    try:
        os.remove(output_file)
    except PermissionError:
        print(f"❌ Cannot overwrite {output_file}. Please close the file if it's open and try again.")
        exit(1)

def verify_required_columns(df: pd.DataFrame) -> Dict[str, str]:
    required_columns = ['memberNumber', 'claimNumber', 'claimTransactionIdentifier']
    missing_columns = [col for col in required_columns if col not in df.columns]
    if missing_columns:
        raise ValueError(f"Missing required columns: {', '.join(missing_columns)}")
    return {col: col for col in required_columns}

def wildcard_path(path):
    return '/'.join(['*' if p.isdigit() else p for p in path.strip('/').split('/')])

def extract_values_with_paths(json_obj, pointer):
    def extract(obj, parts, path_so_far):
        if not parts:
            return [(path_so_far.rstrip('/'), obj)]
        part = parts[0]
        rest = parts[1:]
        results = []
        if isinstance(obj, list):
            if part == '*':
                for idx, item in enumerate(obj):
                    results.extend(extract(item, rest, f"{path_so_far}/{idx}"))
            elif part.isdigit():
                idx = int(part)
                if 0 <= idx < len(obj):
                    results.extend(extract(obj[idx], rest, f"{path_so_far}/{idx}"))
        elif isinstance(obj, dict):
            if part in obj:
                results.extend(extract(obj[part], rest, f"{path_so_far}/{part}"))
        return results
    parts = pointer.strip('/').split('/')
    return extract(json_obj, parts, '')

def safe_float_convert(value):
    try:
        if isinstance(value, str):
            value = value.replace(',', '').replace('+', '')
        return float(value) if value not in ('', None) else 0.0
    except (ValueError, TypeError):
        return 0.0

def normalize_value(v: Any) -> Any:
    return v.strip() if isinstance(v, str) else v

def value_status(a_val: Any, p_val: Any) -> str:
    if a_val is None and p_val is None:
        return "No Value"
    if a_val is None:
        return "Missing in Alex"
    if p_val is None:
        return "Missing in GQL"
    if normalize_value(a_val) == normalize_value(p_val):
        return "Matched"
    return "Different"

def adjusted_values(a_val: Any, p_val: Any, status: str) -> tuple:
    if status == "Missing in Alex":
        return "", p_val
    if status == "Missing in GQL":
        return a_val, ""
    return a_val, p_val

def extract_node_from_path(path: str) -> str:
    return path.strip('/').split('/')[-1] if path else "Unknown"

def extract_canonical(path: str) -> str:
    if isinstance(path, str) and path.startswith("/data/"):
        parts = path.strip("/").split("/")
        if len(parts) > 1:
            return parts[1]
    return "Unknown"

def construct_missing_path(base_path_with_wildcard, found_path_with_index):
    """Constructs a parallel path for a missing item using the index from the found item."""
    match = re.search(r'/(\d+)(?=/|$)', found_path_with_index)
    if match and '*' in base_path_with_wildcard:
        # Replace the first wildcard with the found index
        return base_path_with_wildcard.replace('*', match.group(1), 1)
    return base_path_with_wildcard # Fallback for non-array paths or complex cases

def compare_with_mapping(alex_data: Any, ppkg_data: Any, mapping_path: str, member_number: str, claimTransactionIdentifier: str) -> List[Dict[str, Any]]:
    excel_data = pd.read_excel(mapping_path)
    base_col = excel_data.columns[0]
    target_col = excel_data.columns[1]
    excel_data.dropna(subset=[base_col, target_col], inplace=True)
    mapping = list(zip(excel_data[base_col], excel_data[target_col]))
    rows = []

    alex_claim = alex_data
    ppkg_claim = ppkg_data

    for ppkg_path, alex_path in mapping:
        canonical = extract_canonical(ppkg_path)
        node = extract_node_from_path(ppkg_path)
        ppkg_path_wild = wildcard_path(ppkg_path)

        if '+' in alex_path:
            alex_pointers = [p.strip() for p in alex_path.split('+')]
            alex_vals_total = 0.0
            for pointer in alex_pointers:
                pointer_wild = wildcard_path(pointer)
                vals = extract_values_with_paths(alex_claim, pointer_wild)
                if vals:
                    for _, val in vals:
                        alex_vals_total += safe_float_convert(val)

            ppkg_vals = extract_values_with_paths(ppkg_claim, ppkg_path_wild)
            if not ppkg_vals:
                match_status = 'Missing in GQL' if alex_vals_total != 0 else 'Matched'
                a_out, p_out = adjusted_values(alex_vals_total, None, match_status)
                rows.append({
                    "SubscribeID": member_number, "claimTransactionIdentifier": claimTransactionIdentifier, "Canonical": canonical,
                    "Node": node, "HCP GQL Schema Path": ppkg_path or "Missing",
                    "Custom Resolver Schema Path": alex_path or "Missing", "HCP GQL Value": p_out,
                    "Custom Resolver Value": a_out, "Status": match_status,
                })
            else:
                for p_path_concrete, ppkg_val in ppkg_vals:
                    ppkg_val_float = safe_float_convert(ppkg_val)
                    match_status = 'Matched' if abs(alex_vals_total - ppkg_val_float) < 0.01 else 'Different'
                    a_out, p_out = adjusted_values(alex_vals_total, ppkg_val_float, match_status)
                    rows.append({
                        "SubscribeID": member_number, "claimTransactionIdentifier": claimTransactionIdentifier, "Canonical": canonical,
                        "Node": node, "HCP GQL Schema Path": p_path_concrete,
                        "Custom Resolver Schema Path": alex_path, "HCP GQL Value": p_out,
                        "Custom Resolver Value": a_out, "Status": match_status,
                    })
        else:
            alex_path_wild = wildcard_path(alex_path)
            alex_values = extract_values_with_paths(alex_claim, alex_path_wild)
            ppkg_values = extract_values_with_paths(ppkg_claim, ppkg_path_wild)

            if not alex_values and not ppkg_values:
                rows.append({
                    "SubscribeID": member_number, "claimTransactionIdentifier": claimTransactionIdentifier, "Canonical": canonical,
                    "Node": node, "HCP GQL Schema Path": ppkg_path or "Missing",
                    "Custom Resolver Schema Path": alex_path or "Missing", "HCP GQL Value": "",
                    "Custom Resolver Value": "", "Status": "No Value",
                })
                continue

            alex_map = defaultdict(list)
            for path, val in alex_values:
                base_path = re.sub(r'/\d+(?=/|$)', '/*', path)
                alex_map[base_path].append((path, val))

            ppkg_map = defaultdict(list)
            for path, val in ppkg_values:
                base_path = re.sub(r'/\d+(?=/|$)', '/*', path)
                ppkg_map[base_path].append((path, val))

            all_base_paths = sorted(list(set(alex_map.keys()) | set(ppkg_map.keys())))

            for base_path in all_base_paths:
                alex_list = alex_map.get(base_path, [])
                ppkg_list = ppkg_map.get(base_path, [])

                for p_path_item, p_val in ppkg_list:
                    for a_path_item, a_val in alex_list:
                        status = "Matched" if normalize_value(p_val) == normalize_value(a_val) else "Different"
                        a_out, p_out = adjusted_values(a_val, p_val, status)
                        rows.append({
                            "SubscribeID": member_number, "claimTransactionIdentifier": claimTransactionIdentifier,
                            "Canonical": canonical, "Node": node, "HCP GQL Schema Path": p_path_item,
                            "Custom Resolver Schema Path": a_path_item,
                            "HCP GQL Value": p_out, "Custom Resolver Value": a_out, "Status": status,
                        })

                if not alex_list and ppkg_list:
                    for p_path_item, p_val in ppkg_list:
                        status = "Missing in Alex"
                        a_out, p_out = adjusted_values(None, p_val, status)
                        rows.append({
                            "SubscribeID": member_number, "claimTransactionIdentifier": claimTransactionIdentifier,
                            "Canonical": canonical, "Node": node, "HCP GQL Schema Path": p_path_item,
                            "Custom Resolver Schema Path": alex_path,
                            "HCP GQL Value": p_out, "Custom Resolver Value": a_out, "Status": status,
                        })

                if not ppkg_list and alex_list:
                    for a_path_item, a_val in alex_list:
                        status = "Missing in GQL"
                        a_out, p_out = adjusted_values(a_val, None, status)
                        rows.append({
                            "SubscribeID": member_number, "claimTransactionIdentifier": claimTransactionIdentifier,
                            "Canonical": canonical, "Node": node, "HCP GQL Schema Path": ppkg_path,
                            "Custom Resolver Schema Path": a_path_item,
                            "HCP GQL Value": p_out, "Custom Resolver Value": a_out, "Status": status,
                        })
    return rows

def apply_conditional_formatting(file_path: str):
    wb = load_workbook(file_path)
    ws = wb["Validation Report"]

    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    gray_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

    status_col = None
    for idx, cell in enumerate(ws[1], start=1):
        if cell.value == "Status":
            status_col = idx
            break

    if status_col:
        col_letter = ws.cell(row=1, column=status_col).column_letter
        data_range = f"{col_letter}2:{col_letter}{ws.max_row}"
        ws.conditional_formatting.add(data_range, FormulaRule(formula=[f'${col_letter}2="Matched"'], fill=green_fill))
        ws.conditional_formatting.add(data_range, FormulaRule(formula=[f'${col_letter}2="Different"'], fill=red_fill))
        ws.conditional_formatting.add(data_range, FormulaRule(formula=[f'OR(${col_letter}2="Missing in Alex", ${col_letter}2="Missing in GQL")'], fill=yellow_fill))
        ws.conditional_formatting.add(data_range, FormulaRule(formula=[f'${col_letter}2="No Value"'], fill=gray_fill))

    wb.save(file_path)

def main():
    logging.info("Starting JSON comparison script...")
    all_rows = []

    try:
        df_claims = pd.read_csv(claims_csv_path)
        verify_required_columns(df_claims)
    except FileNotFoundError:
        logging.error(f"Claims file not found at {claims_csv_path}")
        return
    except ValueError as ve:
        logging.error(str(ve))
        return
    except Exception as e:
        logging.error(f"Error reading claims file: {str(e)}")
        return

    for _, row in df_claims.iterrows():
        member_number = str(row.get('memberNumber', '')).zfill(9)
        claim_number = str(row.get('claimNumber', ''))
        claimTransactionIdentifier = str(row.get('claimTransactionIdentifier', ''))
        identifier = f"{member_number}_{claim_number}"

        alex_file = os.path.join(alex_dir, f"{identifier}_alexResponse.json")
        ppkg_file = os.path.join(ppkg_dir, f"{identifier}_GQLResponse.json")

        if not os.path.exists(alex_file):
            logging.warning(f"Missing alex file: {alex_file}")
            continue
        if not os.path.exists(ppkg_file):
            logging.warning(f"Missing GQL file: {ppkg_file}")
            continue

        with open(alex_file, "r", encoding="utf-8") as fa:
            alex_data = json.load(fa)
        with open(ppkg_file, "r", encoding="utf-8") as fp:
            ppkg_data = json.load(fp)

        if os.path.exists(mapping_sheet_path):
            rows = compare_with_mapping(alex_data, ppkg_data, mapping_sheet_path, member_number, claimTransactionIdentifier)
            all_rows.extend(rows)
        else:
            logging.error(f"Mapping sheet not found: {mapping_sheet_path}")
            break

    if all_rows:
        df_report = pd.DataFrame(all_rows)

        # --- Create Summary DataFrame ---
        summary_data = []
        for claim_txn, group in df_report.groupby('claimTransactionIdentifier'):
            relevant_rows = group[group['Status'] != 'No Value']
            total_fields = len(relevant_rows)
            matched_fields = (relevant_rows['Status'] == 'Matched').sum()
            coverage_percentage = (matched_fields / total_fields) * 100 if total_fields > 0 else 0
            summary_data.append({
                'claimTransactionIdentifier': claim_txn,
                'Coverage Percentage': f"{coverage_percentage:.2f}%",
                'Matched Values': f"{matched_fields}/{total_fields}"
            })
        summary_df = pd.DataFrame(summary_data)

        # --- Create Stats By Element DataFrame ---
        stats_data = []
        df_report['BasePath'] = df_report['HCP GQL Schema Path'].str.replace(r'/\d+(?=/|$)', '/*', regex=True)
        for path, group in df_report.groupby('BasePath'):
            relevant_rows = group[group['Status'] != 'No Value']
            total = len(relevant_rows)
            matched = (relevant_rows['Status'] == 'Matched').sum()
            stats_data.append({
                "Field Path": path,
                "# of Occurrences": total,
                "# of Matches": matched,
                "% Matched": (matched / total * 100) if total > 0 else 0.0
            })
        stats_df = pd.DataFrame(stats_data).sort_values(by="% Matched", ascending=False)
        df_report = df_report.drop(columns=['BasePath'])

        # --- Write all DataFrames to a single Excel file ---
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            summary_df.to_excel(writer, sheet_name='Summary', index=False)
            df_report.to_excel(writer, sheet_name="Validation Report", index=False)
            stats_df.to_excel(writer, sheet_name='Stats By Element', index=False)

        apply_conditional_formatting(output_file)
        logging.info(f"✅ Validation report generated with 3 sheets: {output_file}")
    else:
        logging.info("No data to generate report.")

if __name__ == "__main__":
    mapping_sheet_path = os.path.join(dir_path, "src", "test", "utility", "validation_mapping", "alex_ppkg_cosmos", "gql_customResolver", "gql_customresolver.xlsx")
    claims_csv_path = os.path.join(dir_path, 'src', 'test', 'resources', 'testdata', 'gql_alex', 'claimTranscation_TestData.csv')
    main()