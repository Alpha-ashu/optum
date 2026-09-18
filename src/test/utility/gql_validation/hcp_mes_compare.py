import os
import sys
import json
import logging
import re
from typing import Any, List, Dict, Tuple, Optional
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import FormulaRule

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
canonical = sys.argv[1] if len(sys.argv) > 1 else "claimTransaction"
logging.info(f"Consumer argument received: {canonical}")

dir_path = os.getcwd().split("src")[0]

hcp_dir = os.path.join(dir_path, 'target', 'All_Responses', 'HCP_Responses')
mes_dir = os.path.join(dir_path, 'target', 'All_Responses', 'ALEX_Responses')
output_dir = os.path.join(dir_path, 'target', 'All_Responses', 'Filtered_Responses')
os.makedirs(output_dir, exist_ok=True)
output_file = os.path.join(output_dir, f"{canonical}_validated_report.xlsx")

if os.path.exists(output_file):
    try:
        os.remove(output_file)
    except PermissionError:
        print(f"❌ Cannot overwrite {output_file}. Please close the file if it's open and try again.")
        sys.exit(1)

def type_name(value: Any) -> str:
    """Return a human-readable type name for JSON-like values."""
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "number"
    if isinstance(value, str):
        return "string"
    if isinstance(value, list):
        return "array"
    if isinstance(value, dict):
        return "object"
    return type(value).__name__

def extract_canonical(path: str) -> str:
    if isinstance(path, str) and path.startswith("/data/"):
        parts = path.strip("/").split("/")
        if len(parts) > 1:
            return parts[1]
    return "Unknown"

def extract_node_from_path(path: str) -> str:
    return path.strip('/').split('/')[-1] if path else "Unknown"

def extract_cti_from_json(json_data: Any) -> Optional[str]:
    try:
        if isinstance(json_data, dict):
            data = json_data.get('data')
            if isinstance(data, dict):
                cts = data.get('claimTransactions')
                if isinstance(cts, list) and len(cts) > 0:
                    first_claim = cts[0]
                    if isinstance(first_claim, dict):
                        cti = first_claim.get('claimTransactionIdentifier')
                        if cti is not None:
                            return str(cti)
    except Exception as e:
        logging.debug(f"Error extracting CTI from JSON: {e}")
    return None

def extract_cti_from_filename(filename: str) -> Optional[str]:
    try:
        parts = filename.replace('_HCPResponse.json', '').replace('_ALEXResponse.json', '').split('_')
        if len(parts) >= 2:
            return parts[1]
    except Exception as e:
        logging.debug(f"Error extracting CTI from filename {filename}: {e}")
    return None

def discover_files() -> Tuple[Dict[str, Dict], Dict[str, Dict]]:
    hcp_map: Dict[str, Dict[str, Any]] = {}
    mes_map: Dict[str, Dict[str, Any]] = {}

    if os.path.exists(hcp_dir):
        for filename in os.listdir(hcp_dir):
            if filename.endswith('.json'):
                file_path = os.path.join(hcp_dir, filename)
                try:
                    with open(file_path, 'r', encoding='utf-8') as f:
                        data = json.load(f)
                        cti = extract_cti_from_json(data) or extract_cti_from_filename(filename)
                        if cti:
                            hcp_map[cti] = {'file': filename, 'data': data}
                            logging.debug(f"HCP: {filename} → CTI: {cti}")
                        else:
                            logging.warning(f"Could not extract CTI from HCP file: {filename}")
                except Exception as e:
                    logging.error(f"Error reading HCP file {filename}: {e}")

    if os.path.exists(mes_dir):
        for filename in os.listdir(mes_dir):
            if filename.endswith('.json'):
                file_path = os.path.join(mes_dir, filename)
                try:
                    with open(file_path, 'r', encoding='utf-8') as f:
                        data = json.load(f)
                        cti = extract_cti_from_json(data) or extract_cti_from_filename(filename)
                        if cti:
                            mes_map[cti] = {'file': filename, 'data': data}
                            logging.debug(f"ALEX: {filename} → CTI: {cti}")
                        else:
                            logging.warning(f"Could not extract CTI from ALEX file: {filename}")
                except Exception as e:
                    logging.error(f"Error reading ALEX file {filename}: {e}")

    logging.info(f"Discovered {len(hcp_map)} HCP files and {len(mes_map)} ALEX files")
    return hcp_map, mes_map

def get_value_at_path(json_obj: Any, path: str) -> Tuple[bool, Any]:
    if not path:
        return False, None
    parts = path.strip('/').split('/')
    current: Any = json_obj
    for part in parts:
        if current is None:
            return True, None
        if isinstance(current, dict):
            if part in current:
                current = current[part]
            else:
                return False, None
        elif isinstance(current, list):
            if part.isdigit():
                idx = int(part)
                if 0 <= idx < len(current):
                    current = current[idx]
                else:
                    return False, None
            else:
                return False, None
        else:
            return False, None
    return True, current

def expand_hcp_iterators(json_data: Any, path: str) -> List[str]:
    if '*' not in path:
        return [path]
    parts = path.strip('/').split('/')
    expanded = ['']
    for part in parts:
        if part == '*':
            new_expanded: List[str] = []
            for prefix in expanded:
                exists, current = get_value_at_path(json_data, prefix)
                if exists and isinstance(current, list):
                    for i in range(len(current)):
                        new_expanded.append(f"{prefix}/{i}")
                else:
                    new_expanded.append(f"{prefix}/*")
            expanded = new_expanded
        else:
            expanded = [f"{prefix}/{part}" for prefix in expanded]
    return [p for p in expanded if '*' not in p]

def substitute_indices(mes_path: str, hcp_concrete_path: str, hcp_template_path: str) -> str:
    if '*' not in mes_path:
        return mes_path
    hcp_parts = hcp_concrete_path.strip('/').split('/')
    hcp_template_parts = hcp_template_path.strip('/').split('/')
    mes_parts = mes_path.strip('/').split('/')
    index_map: Dict[int, str] = {}
    wildcard_count = 0
    for concrete, template in zip(hcp_parts, hcp_template_parts):
        if template == '*' and concrete.isdigit():
            index_map[wildcard_count] = concrete
            wildcard_count += 1
    result_parts: List[str] = []
    wildcard_count = 0
    for part in mes_parts:
        if part == '*':
            result_parts.append(index_map.get(wildcard_count, '*'))
            wildcard_count += 1
        else:
            result_parts.append(part)
    return '/' + '/'.join(result_parts)

def strict_equals(val1: Any, val2: Any) -> bool:
    if type(val1) != type(val2):
        return False
    if val1 is None and val2 is None:
        return True
    return val1 == val2

def classify_status(hcp_exists: bool, hcp_val: Any, mes_exists: bool, mes_val: Any) -> str:
    if hcp_exists and hcp_val is None:
        return "Null Value"
    if mes_exists and mes_val is None:
        return "Null Value"
    if not hcp_exists and not mes_exists:
        return "No Value"
    if not hcp_exists:
        return "Missing in HCP"
    if not mes_exists:
        return "Missing in ALEX"
    if strict_equals(hcp_val, mes_val):
        return "Matched"
    return "Different"

def load_mapping_sheet(mapping_path: str) -> pd.DataFrame:
    df = pd.read_excel(mapping_path)
    hcp_col = None
    mes_col = None
    canonical_col = None
    for col in df.columns:
        col_lower = str(col).lower()
        if 'hcp' in col_lower and 'path' in col_lower:
            hcp_col = col
        elif 'mes' in col_lower and 'path' in col_lower:
            mes_col = col
        elif 'canonical' in col_lower:
            canonical_col = col
    if not hcp_col or not mes_col:
        hcp_col = df.columns[0]
        mes_col = df.columns[1]
    result = df[[hcp_col, mes_col]].copy()
    result.columns = ['HCP_Path', 'ALEX_Path']
    if canonical_col:
        result['Canonical'] = df[canonical_col]
    else:
        result['Canonical'] = ''
    result.dropna(subset=['HCP_Path', 'ALEX_Path'], inplace=True)
    return result

def compare_single_cti(cti: str, hcp_data: Any, mes_data: Any, mapping_df: pd.DataFrame) -> List[Dict]:
    rows: List[Dict[str, Any]] = []
    for _, map_row in mapping_df.iterrows():
        hcp_template = str(map_row['HCP_Path']).strip()
        mes_template = str(map_row['ALEX_Path']).strip()
        canonical_from_path = extract_canonical(hcp_template)
        canonical_val = str(map_row['Canonical']) if map_row['Canonical'] else canonical_from_path
        hcp_concrete_paths = expand_hcp_iterators(hcp_data, hcp_template)
        for hcp_concrete in hcp_concrete_paths:
            if '*' in hcp_concrete:
                logging.debug(f"Skipping unexpanded HCP path: {hcp_concrete}")
                continue
            mes_concrete = substitute_indices(mes_template, hcp_concrete, hcp_template)
            if '*' in mes_concrete:
                logging.debug(f"Skipping ALEX path with unresolved wildcards: {mes_concrete}")
                continue
            node = extract_node_from_path(hcp_concrete)
            hcp_exists, hcp_val = get_value_at_path(hcp_data, hcp_concrete)
            mes_exists, mes_val = get_value_at_path(mes_data, mes_concrete)
            status = classify_status(hcp_exists, hcp_val, mes_exists, mes_val)
            rows.append({
                'claimTransactionIdentifier': cti,
                'Canonical': canonical_val,
                'Node': node,
                'HCP Path': hcp_concrete,
                'ALEX Path': mes_concrete,
                'HCP Value': hcp_val if hcp_exists else '',
                'ALEX Value': mes_val if mes_exists else '',
                #'HCP Type': type_name(hcp_val) if hcp_exists else '',
                #'ALEX Type': type_name(mes_val) if mes_exists else '',
                #'Null Value': 'Yes' if status == 'Null Value' else 'No',
                'Status': status
            })
    return rows

def build_response_schema(df_report: pd.DataFrame) -> pd.DataFrame:
    def normalize_path(p: str) -> str:
        if not isinstance(p, str):
            return ''
        return re.sub(r'/\d+(?=/|$)', '/*', p)
    schema_rows: List[Dict[str, Any]] = []
    grouped = df_report.groupby(['HCP Path', 'ALEX Path'], dropna=False)
    for (hcp_path, mes_path), group in grouped:
        hcp_norm = normalize_path(hcp_path)
        mes_norm = normalize_path(mes_path)
        hcp_present_mask = ~group['Status'].isin(['Missing in HCP', 'No Value'])
        mes_present_mask = ~group['Status'].isin(['Missing in ALEX', 'No Value'])
        hcp_types = group['HCP Type'].replace('', pd.NA).dropna()
        mes_types = group['ALEX Type'].replace('', pd.NA).dropna()
        schema_rows.append({
            'Canonical': group.iloc[0]['Canonical'] if len(group) > 0 else '',
            'HCP Field Path': hcp_norm,
            'ALEX Field Path': alex_norm,
            'HCP Types Observed': ', '.join(sorted(set(hcp_types.tolist()))) if len(hcp_types) > 0 else '',
            'ALEX Types Observed': ', '.join(sorted(set(mes_types.tolist()))) if len(mes_types) > 0 else '',
            'HCP Present Count': int(hcp_present_mask.sum()),
            'ALEX Present Count': int(mes_present_mask.sum()),
            'Total Rows': int(len(group))
        })
    schema_df = pd.DataFrame(schema_rows).drop_duplicates()
    return schema_df.sort_values(by=['Canonical', 'HCP Field Path', 'ALEX Field Path'])
def apply_conditional_formatting(file_path: str):
    wb = load_workbook(file_path)
    if 'Validation Report' not in wb.sheetnames:
        wb.save(file_path)
        return
    ws = wb['Validation Report']
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    gray_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    orange_fill = PatternFill(start_color="FFD9B3", end_color="FFD9B3", fill_type="solid")
    status_col = None
    for idx, cell in enumerate(ws[1], start=1):
        if cell.value == "Status":
            status_col = idx
            break
    if status_col:
        col_letter = ws.cell(row=1, column=status_col).column_letter
        data_range = f"{col_letter}2:{col_letter}{ws.max_row}"
        ws.conditional_formatting.add(data_range, FormulaRule(formula=[f'${col_letter}2="Matched"'], fill=green_fill))
        ws.conditional_formatting.add(data_range, FormulaRule(formula=[f'${col_letter}2="Different"'], fill=red_fill))
        ws.conditional_formatting.add(data_range, FormulaRule(formula=[f'OR(${col_letter}2="Missing in ALEX", ${col_letter}2="Missing in HCP")'], fill=yellow_fill))
        ws.conditional_formatting.add(data_range, FormulaRule(formula=[f'${col_letter}2="No Value"'], fill=gray_fill))
        ws.conditional_formatting.add(data_range, FormulaRule(formula=[f'${col_letter}2="Null Value"'], fill=orange_fill))
    wb.save(file_path)

def main():
    logging.info("Starting strict iterator-based validation with GENERIC array handling...")
    # Determine mapping sheet path
    if canonical == "claimTransaction":
        mapping_sheet_path = os.path.join(dir_path, "src", "test", "utility", "validation_mapping", "alex_ppkg_cosmos", "hcp_mes", "hcp_mes_claimTransaction.xlsx")
    elif canonical == "Claimservicelines":
        mapping_sheet_path = os.path.join(dir_path, "src", "test", "utility", "validation_mapping", "alex_ppkg_cosmos", "hcp_mes", "hcp_mes_claimservicelines.xlsx")
    elif canonical == "claimAccumulators":
        mapping_sheet_path = os.path.join(dir_path, "src", "test", "utility", "validation_mapping", "alex_ppkg_cosmos", "hcp_mes", "hcp_mes_claimAccumulators.xlsx")
    elif canonical == "claimAmbulances":
        mapping_sheet_path = os.path.join(dir_path, "src", "test", "utility", "validation_mapping", "alex_ppkg_cosmos", "hcp_mes", "hcp_mes_claimAmbulances.xlsx")
    elif canonical == "claimOtherPayers":
        mapping_sheet_path = os.path.join(dir_path, "src", "test", "utility", "validation_mapping", "alex_ppkg_cosmos", "hcp_mes", "hcp_mes_claimOtherPayers.xlsx")
    elif canonical == "claimPricings":
        mapping_sheet_path = os.path.join(dir_path, "src", "test", "utility", "validation_mapping", "alex_ppkg_cosmos", "hcp_mes", "hcp_mes_claimPricings.xlsx")
    elif canonical == "claimServiceAccumulators":
        mapping_sheet_path = os.path.join(dir_path, "src", "test", "utility", "validation_mapping", "alex_ppkg_cosmos", "hcp_mes", "hcp_mes_claimServiceAccumulator.xlsx")
    elif canonical == "claimServiceOtherPayers":
        mapping_sheet_path = os.path.join(dir_path, "src", "test", "utility", "validation_mapping", "alex_ppkg_cosmos", "hcp_mes", "hcp_mes_claimServiceOtherPayers.xlsx")
    elif canonical == "claimServicePricings":
        mapping_sheet_path = os.path.join(dir_path, "src", "test", "utility", "validation_mapping", "alex_ppkg_cosmos", "hcp_mes", "hcp_mes_claimServicePricings.xlsx")
    elif canonical == "serviceAmbulances":
        mapping_sheet_path = os.path.join(dir_path, "src", "test", "utility", "validation_mapping", "alex_ppkg_cosmos", "hcp_mes", "hcp_mes_serviceAmbulances.xlsx")
    else:
        mapping_sheet_path = os.path.join(dir_path, "src", "test", "utility", "validation_mapping", "alex_ppkg_cosmos", "`hcp_gql`", "hcp_gql_claimTransaction.xlsx")

    if not os.path.exists(mapping_sheet_path):
        logging.error(f"Mapping sheet not found: {mapping_sheet_path}")
        return
    try:
        mapping_df = load_mapping_sheet(mapping_sheet_path)
        logging.info(f"Loaded {len(mapping_df)} mapping rows")
    except Exception as e:
        logging.error(f"Error loading mapping sheet: {e}")
        return
    hcp_map, mes_map = discover_files()
    common_ctis = set(hcp_map.keys()) & set(mes_map.keys())
    logging.info(f"Found {len(common_ctis)} common CTIs to validate")
    if not common_ctis:
        logging.warning("No common CTIs found between HCP and ALEX")
        return
    all_rows: List[Dict[str, Any]] = []
    for cti in sorted(common_ctis):
        logging.info(f"Validating CTI: {cti}")
        hcp_data = hcp_map[cti]['data']
        mes_data = mes_map[cti]['data']
        rows = compare_single_cti(cti, hcp_data, mes_data, mapping_df)
        all_rows.extend(rows)
        logging.info(f"  → Generated {len(rows)} validation rows for CTI {cti}")
    if not all_rows:
        logging.info("No validation data generated")
        return
    df_report = pd.DataFrame(all_rows)
    summary_data: List[Dict[str, Any]] = []
    for cti, group in df_report.groupby('claimTransactionIdentifier'):
        canonical_val = group.iloc[0]['Canonical'] if len(group) > 0 else canonical
        matched_count = int((group['Status'] == 'Matched').sum())
        different_count = int((group['Status'] == 'Different').sum())
        missing_hcp_count = int((group['Status'] == 'Missing in HCP').sum())
        missing_mes_count = int((group['Status'] == 'Missing in ALEX').sum())
        no_value_count = int((group['Status'] == 'No Value').sum())
        null_value_count = int((group['Status'] == 'Null Value').sum())
        mapping_errors = 0
        relevant_total = matched_count + different_count + missing_hcp_count + missing_mes_count
        coverage = (matched_count / relevant_total * 100) if relevant_total > 0 else 0
        summary_data.append({
            'Canonical': canonical_val,
            'claimTransactionIdentifier': cti,
            'Matched (exact)': matched_count,
            'Different': different_count,
            'Missing in HCP': missing_hcp_count,
            'Missing in ALEX': missing_mes_count,
            'No Value': no_value_count,
            'Null Value': null_value_count,
            'Mapping Errors': mapping_errors,
            'Relevant Total (excl No Value & Null)': relevant_total,
            'Coverage % (exact)': round(coverage, 2)
        })
    summary_df = pd.DataFrame(summary_data)
    stats_data: List[Dict[str, Any]] = []
    for (hcp_path, mes_path), group in df_report.groupby(['HCP Path', 'ALEX Path']):
        hcp_normalized = re.sub(r'/\d+(?=/|$)', '/*', hcp_path)
        mes_normalized = re.sub(r'/\d+(?=/|$)', '/*', mes_path)
        relevant = group[(group['Status'] != 'No Value') & (group['Status'] != 'Null Value')]
        total = int(len(relevant))
        matched = int((relevant['Status'] == 'Matched').sum())
        stats_data.append({
            'HCP Field Path': hcp_normalized,
            'ALEX Field Path': mes_normalized,
            '# of Occurrences': total,
            '# of Matches': matched,
            '% Matched': round((matched / total * 100), 2) if total > 0 else 0.0
        })
    stats_df = pd.DataFrame(stats_data).drop_duplicates().sort_values(by='% Matched', ascending=False)
    schema_df = build_response_schema(df_report)
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        summary_df.to_excel(writer, sheet_name='Summary', index=False)
        df_report.to_excel(writer, sheet_name='Validation Report', index=False)
        stats_df.to_excel(writer, sheet_name='Stats By Element', index=False)
        schema_df.to_excel(writer, sheet_name='Response Schema', index=False)
    apply_conditional_formatting(output_file)
    logging.info(f"✅ Validation report generated: {output_file}")
    logging.info(f"   Total validation rows: {len(all_rows)}")

if __name__ == "__main__":
    main()