"""
Generic, Configuration-Driven Validation Framework
==================================================
This module turns the (unchanged) comparison + report-generation engine in
``Schema_Validation.py`` into a **generic**, configuration-driven framework.

Goal
----
Onboard a new API validation (Clink Summary, Clink Physician, Clink Hospital,
Claims360 Summary, Claims360 Detail, future APIs, …) WITHOUT touching any
Python comparison / validation / mapping / report logic. To add a validation a
user only provides:

    validation_config/<Name>/config.json     (+ optional feature file & mapping)

This module is responsible ONLY for *discovery* and *wiring* — it does NOT
implement any field comparison, validation rules, field mapping, or Excel
report generation. Those remain exactly as they are in ``Schema_Validation.py``.

Responsibilities
----------------
1. Discover every validation defined under ``validation_config/``.
2. Parse the Karate feature file to auto-extract metadata:
      Feature Name, Module, Type, Scenario Name, API URL, Environment.
3. Determine the Validation Type (Summary / Physician / Hospital / Detail …)
   from config, feature metadata, or folder structure — no hardcoding.
4. Dynamically detect response files per source (PPKG / Alex / UPM / Claim360 /
   Clink / future) by scanning folders and matching a common claim key.
5. Build the CONSUMER_CONFIG entry consumed by the existing ``run()``.
6. Build a dynamic report name and a report-header metadata block.

Nothing here mutates the behaviour of the existing three consumers: every
config key it introduces is optional and defaults to the current behaviour.
"""

from __future__ import annotations

import os
import re
import json
import glob
from datetime import datetime
from typing import Dict, List, Optional, Tuple


# ─────────────────────────────────────────────────────────────────────────────
# Project base directory (mirrors alex_hcp_compare.BASE_DIR resolution)
# ─────────────────────────────────────────────────────────────────────────────
BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..', '..'))
if not os.path.exists(os.path.join(BASE_DIR, 'pom.xml')):
    BASE_DIR = os.getcwd().split("src")[0] if "src" in os.getcwd() else os.getcwd()

# Location of the configuration repository. Overridable via env var so the
# framework can point at a shared/external config store without code changes.
VALIDATION_CONFIG_DIR = os.environ.get(
    'VALIDATION_CONFIG_DIR',
    os.path.join(os.path.dirname(__file__), 'validation_config'),
)

# Canonical validation types recognised from feature text / folder names.
KNOWN_VALIDATION_TYPES = ['Summary', 'Physician', 'Hospital', 'Detail', 'Professional', 'Institutional']

# Default response-source definitions. A source maps a logical name (used in the
# report header) to the folder + filename suffix that holds its JSON responses.
# The engine compares exactly two sources: an "expected" and an "actual" one.
DEFAULT_SOURCES: Dict[str, Dict[str, str]] = {
    'PPKG':     {'folder': 'PPKG_Responses',     'suffix': '_ppkgResponse.json'},
    'Alex':     {'folder': 'Alex_Responses',     'suffix': '_alexResponse.json'},
    'UPM':      {'folder': 'UPM_Responses',      'suffix': '_upmResponse.json'},
    'Claim360': {'folder': 'Claim360_Responses', 'suffix': '_claim360Response.json'},
    'Clink':    {'folder': 'Clink_Responses',    'suffix': '_clinkResponse.json'},
}


# ─────────────────────────────────────────────────────────────────────────────
# Path helpers
# ─────────────────────────────────────────────────────────────────────────────

def _abs(path: Optional[str]) -> Optional[str]:
    """Resolve a possibly-relative path against BASE_DIR."""
    if not path:
        return path
    return path if os.path.isabs(path) else os.path.join(BASE_DIR, path)


# ─────────────────────────────────────────────────────────────────────────────
# 1. Feature file metadata extraction
# ─────────────────────────────────────────────────────────────────────────────

def parse_feature_file(feature_path: Optional[str]) -> Dict[str, str]:
    """
    Parse a Karate ``.feature`` file and extract metadata.

    Returns a dict (any/all keys optional depending on the file):
        feature_name, module, type, scenario_name, api_url, environment

    Best-effort and fully tolerant — a missing file or unpar-seable line simply
    yields an empty / partial dict, never an exception.
    """
    meta: Dict[str, str] = {}
    if not feature_path:
        return meta
    feature_path = _abs(feature_path)
    if not feature_path or not os.path.exists(feature_path):
        return meta

    try:
        with open(feature_path, 'r', encoding='utf-8', errors='replace') as f:
            text = f.read()
    except OSError:
        return meta

    lines = text.splitlines()

    # ── Feature:  e.g. "Feature: Clink API Payer - Claims 360 Summary Search"
    for ln in lines:
        m = re.match(r'\s*Feature\s*:\s*(.+)', ln, re.IGNORECASE)
        if m:
            full = m.group(1).strip()
            meta['feature_name'] = full
            module, ftype = _split_feature_title(full)
            if module:
                meta['module'] = module
            if ftype:
                meta['type'] = ftype
            break

    # ── Scenario / Scenario Outline (may span the next non-empty line) ──
    for i, ln in enumerate(lines):
        m = re.match(r'\s*Scenario(?:\s+Outline)?\s*:\s*(.*)', ln, re.IGNORECASE)
        if m:
            scen = m.group(1).strip()
            if not scen:  # title sits on the following line
                for nxt in lines[i + 1:]:
                    if nxt.strip():
                        scen = nxt.strip()
                        break
            meta['scenario_name'] = _clean_scenario(scen)
            break

    # ── API URL:  * def api_url = 'https://…'   (single or double quotes) ──
    m = re.search(r"""def\s+api_url\s*=\s*['"]([^'"]+)['"]""", text, re.IGNORECASE)
    if not m:  # any http(s) url as a fallback
        m = re.search(r"""['"](https?://[^'"]+)['"]""", text)
    if m:
        meta['api_url'] = m.group(1).strip()

    # ── Environment: explicit def, else inferred from URL / feature text ──
    env = _extract_environment(text, meta.get('api_url', ''))
    if env:
        meta['environment'] = env

    return meta


def _split_feature_title(full: str) -> Tuple[str, str]:
    """
    Split a feature title such as
        "Clink API Payer - Claims 360 Summary Search"
    into  (module, type)  →  ("Claims360", "Summary").

    The part before the first ' - ' is treated as the feature/consumer label;
    the remainder is scanned for the module (e.g. "Claims 360" → "Claims360")
    and a known validation type keyword.
    """
    module = ''
    ftype = ''
    remainder = full
    if ' - ' in full:
        remainder = full.split(' - ', 1)[1]

    # Module: collapse "Claims 360" → "Claims360", "Claim 360" → "Claim360"
    mm = re.search(r'(Claims?\s*360|Clink|PPKG|Alex|UPM)', remainder, re.IGNORECASE)
    if mm:
        module = re.sub(r'\s+', '', mm.group(1))

    for kw in KNOWN_VALIDATION_TYPES:
        if re.search(rf'\b{kw}\b', remainder, re.IGNORECASE):
            ftype = kw
            break
    return module, ftype


def _clean_scenario(scen: str) -> str:
    """Normalise a scenario title (drop trailing punctuation, tidy spacing)."""
    scen = re.sub(r'\s+', ' ', scen).strip()
    scen = scen.strip('-:').strip()
    return scen


# Canonical environment labels for known upstream/environment tokens.
ENV_LABELS: Dict[str, str] = {
    'qae': 'QAE', 'uat': 'UAT', 'prod': 'PROD', 'production': 'PROD',
    'stage': 'STAGE', 'staging': 'STAGE', 'test': 'Test', 'tst': 'Test',
    'dev': 'DEV', 'development': 'Development', 'nonprod': 'Nonprod',
    'non-prod': 'Nonprod', 'np': 'Nonprod',
    'de-canary': 'Decanary', 'decanary': 'Decanary', 'canary': 'Canary',
    'test': 'Test', 'tst': 'Test', 'demo': 'Demo', 'sandbox': 'Sandbox',
}


def _label_env(value: str) -> str:
    """Normalise a raw environment token into a friendly display label."""
    v = (value or '').strip()
    if not v:
        return ''
    return ENV_LABELS.get(v.lower(), v if len(v) <= 4 else v.title())


def _extract_environment(text: str, api_url: str) -> str:
    """Infer the environment from the feature file.

    Priority:
      1. The ``X-Upstream-Env`` request header (the reliable, explicit signal).
      2. An explicit ``* def env = '...'`` / ``def environment = '...'`` OR a
         Karate ``header environment = '...'`` step (e.g. de-canary routing).
      3. Keyword heuristics on the URL / feature text.
    """
    # 1. X-Upstream-Env header — the authoritative environment indicator.
    m = re.search(r"""X-Upstream-Env\s*=\s*['"]([^'"]+)['"]""", text, re.IGNORECASE)
    if m:
        return _label_env(m.group(1))

    # 2. Explicit def env / environment OR a Karate `header environment = '...'`
    #    step. The `header environment = 'de-canary'` line is how the target
    #    feature routes to the de-canary stack, so it must be recognised here.
    m = re.search(
        r"""(?:def|header)\s+(?:env|environment)\s*=\s*['"]([^'"]+)['"]""",
        text, re.IGNORECASE,
    )
    if m:
        return _label_env(m.group(1))

    # 3. Keyword heuristics.
    haystack = f"{api_url} {text}".lower()
    env_map = [
        ('de-canary', 'Decanary'), ('decanary', 'Decanary'),
        ('nonprod', 'Nonprod'), ('non-prod', 'Nonprod'),
        ('qae', 'QAE'), ('stage', 'STAGE'), ('staging', 'STAGE'),
        ('uat', 'UAT'), ('prod', 'PROD'), ('/dev/', 'DEV'),
    ]
    for needle, label in env_map:
        if needle in haystack:
            return label
    return ''


def build_source_target_metadata(config: Dict) -> Dict[str, str]:
    """Parse the SOURCE and TARGET feature files and return their URL + env.

    ``source_*`` describes the expected system (PPKG / HCP), ``target_*`` the
    actual system (Alex). Explicit config values always win over parsed ones so
    a validation can override auto-detection when needed.
    """
    src = parse_feature_file(config.get('source_feature_file'))
    tgt = parse_feature_file(config.get('target_feature_file'))
    return {
        'source_url':          config.get('source_url')          or src.get('api_url', ''),
        'target_url':          config.get('target_url')          or tgt.get('api_url', ''),
        'source_env':          config.get('source_environment')  or src.get('environment', ''),
        'target_env':          config.get('target_environment')  or tgt.get('environment', ''),
        'source_feature_name': src.get('feature_name', ''),
        'target_feature_name': tgt.get('feature_name', ''),
    }


# ─────────────────────────────────────────────────────────────────────────────
# 2. Validation-type determination (no hardcoding)
# ─────────────────────────────────────────────────────────────────────────────

def determine_validation_type(config: Dict, feature_meta: Dict[str, str]) -> str:
    """
    Resolve the validation type, in priority order:
      1. explicit ``validation_type`` in config
      2. ``response_type`` in config
      3. ``type`` parsed from the feature title
      4. a known type keyword found in the response_base / folder structure
      5. ``claim_type`` in config (legacy)
    """
    for key in ('validation_type', 'response_type'):
        if config.get(key):
            return str(config[key]).strip().title()

    if feature_meta.get('type'):
        return feature_meta['type'].title()

    # Infer from folder structure (…/summary/… , …/physician/… , …)
    for candidate in (config.get('response_base'), config.get('feature_file')):
        if not candidate:
            continue
        low = str(candidate).lower()
        for kw in KNOWN_VALIDATION_TYPES:
            if kw.lower() in low:
                return kw

    return str(config.get('claim_type', '')).strip().title() or 'Summary'


# ─────────────────────────────────────────────────────────────────────────────
# 3. Config repository discovery
# ─────────────────────────────────────────────────────────────────────────────

def discover_validation_configs() -> Dict[str, dict]:
    """
    Scan ``validation_config/`` for ``*.json`` files and build a
    CONSUMER_CONFIG-compatible entry for each. The returned dict is keyed by the
    validation's registration name (config ``name`` or the file name) and can
    be merged straight into ``alex_hcp_compare.CONSUMER_CONFIG``.

    A folder may hold EITHER a single ``config.json`` OR several differently
    named ``*.json`` files (e.g. ``validation_config/IIM/Summary.json``,
    ``Physician.json``, ``Physician_Prod.json``) — both layouts work
    identically, so related validations can share one folder instead of each
    needing its own sub-directory.
    """
    registry: Dict[str, dict] = {}
    if not os.path.isdir(VALIDATION_CONFIG_DIR):
        return registry

    for cfg_path in glob.glob(os.path.join(VALIDATION_CONFIG_DIR, '*', '*.json')):
        # Skip helper/template folders (names starting with '_' or '.') so
        # copy-paste starting points like validation_config/_TEMPLATE/ are never
        # registered as real, runnable validations.
        folder_name = os.path.basename(os.path.dirname(cfg_path))
        if folder_name.startswith(('_', '.')):
            continue
        try:
            entry = build_consumer_entry(cfg_path)
        except Exception as exc:  # never let one bad config break discovery
            print(f"⚠️  Skipping invalid config '{cfg_path}': {exc}")
            continue
        if entry:
            name = entry.pop('__register_name__')
            registry[name] = entry
    return registry


def build_consumer_entry(cfg_path: str) -> Optional[dict]:
    """
    Read a single ``config.json`` and turn it into a CONSUMER_CONFIG entry
    consumed by the existing ``run()``. Auto-resolves the mapping sheet from the
    config folder when not explicitly listed, and enriches the entry with
    feature-file metadata + generic source / report keys.
    """
    with open(cfg_path, 'r', encoding='utf-8') as f:
        config = json.load(f)
    # Fall back name: the FILE name (without extension), not the folder name -
    # several differently-named *.json files can share one folder, so a
    # folder-based fallback would make every un-named config collide.
    file_stem = os.path.splitext(os.path.basename(cfg_path))[0]
    return build_consumer_entry_from_config(
        config, os.path.dirname(cfg_path), default_name=file_stem,
    )


def build_consumer_entry_from_config(
    config: dict, folder: Optional[str] = None, default_name: Optional[str] = None,
) -> Optional[dict]:
    """
    Turn an already-loaded config dict into a CONSUMER_CONFIG entry. Shared by
    both the on-disk discovery (:func:`build_consumer_entry`) and the ad-hoc CLI
    mode, so a validation can be launched from a JSON file OR straight from
    command-line arguments with identical behaviour.

    ``folder`` is where relative mapping sheets are looked up when
    ``mapping_sheets`` is not given (defaults to the project base directory).
    ``default_name`` is the fallback registration name when the config has
    neither ``name`` nor ``validation_name`` (defaults to the folder name).
    """
    folder = folder or BASE_DIR
    folder_name = os.path.basename(folder.rstrip(os.sep)) if folder else ''
    register_name = config.get('name') or config.get('validation_name') or default_name or folder_name

    # ── Parse the feature file for metadata (URL, env, feature, scenario) ──
    feature_meta = parse_feature_file(config.get('feature_file'))

    # ── Validation type & sources ──
    validation_type = determine_validation_type(config, feature_meta)
    expected_source = config.get('expected_source', 'PPKG')
    actual_source   = config.get('actual_source', 'Alex')

    # Merge user-supplied source overrides on top of the sensible defaults.
    sources = {k: dict(v) for k, v in DEFAULT_SOURCES.items()}
    for name, override in (config.get('sources') or {}).items():
        sources.setdefault(name, {})
        sources[name].update(override)

    # ── Mapping sheets: explicit list, else pattern-derived from
    # mapping_sheet_dir + mapping_sheet_pattern, else any .xlsx next to config.
    # The pattern form lets an entire family of validations that share ONE
    # mapping-sheet-naming convention (e.g. every GQL "canonical" type living in
    # json_mapping/ppkg_alex/ppkg_mes_gql/hcp_mes_<name>.xlsx) reuse the SAME
    # engine with zero code changes - only a small per-canonical config.json.
    mapping_sheets = [_abs(m) for m in (config.get('mapping_sheets') or [])]
    if not mapping_sheets:
        ms_dir = config.get('mapping_sheet_dir')
        if ms_dir:
            pattern  = config.get('mapping_sheet_pattern', '{name}.xlsx')
            fname    = pattern.format(name=register_name, canonical=register_name)
            candidate = _abs(os.path.join(ms_dir, fname))
            if os.path.exists(candidate):
                mapping_sheets = [candidate]
            else:
                print(f"⚠️  mapping_sheet_dir pattern resolved to a missing file: {candidate}")
    if not mapping_sheets:
        mapping_sheets = sorted(glob.glob(os.path.join(folder, '*.xlsx')))

    # ── Response root: explicit override, else target/All_Responses/<type> ──
    response_base = config.get('response_base')
    response_base = _abs(response_base) if response_base else None

    entry = {
        '__register_name__': register_name,
        # Keys consumed by the existing run() ---------------------------------
        'test_data':      _abs(config.get('test_data')),
        'claim_filter':   config.get('claim_filter'),
        'claim_type':     config.get('claim_type') or validation_type,
        'mapping_sheets': mapping_sheets,
        # Generic framework extensions (all optional / backward-compatible) ----
        'validation_type': validation_type,
        'expected_source': expected_source,
        'actual_source':   actual_source,
        'sources':         sources,
        'response_base':   response_base,
        'feature_file':    _abs(config.get('feature_file')),
        'config_dir':      folder,
        'metadata': {
            'validation_name': config.get('validation_name') or register_name,
            'feature_name':    config.get('feature_name') or feature_meta.get('feature_name', ''),
            'module':          feature_meta.get('module', ''),
            'scenario_name':   config.get('scenario_name') or feature_meta.get('scenario_name', ''),
            'environment':     config.get('environment') or feature_meta.get('environment', ''),
            'api_url':         config.get('api_url') or feature_meta.get('api_url', ''),
            'validation_type': validation_type,
            'expected_source': expected_source,
            'actual_source':   actual_source,
        },
    }
    # Pass through every remaining config.json key verbatim so new,
    # engine-level options (comparison_mode, mapping_columns, search_type,
    # source_feature_file, ...) flow through WITHOUT the framework needing to
    # know about them. Existing (already-resolved / absolutized) keys win.
    for _k, _v in config.items():
        entry.setdefault(_k, _v)
    return entry


# ─────────────────────────────────────────────────────────────────────────────
# 4. Dynamic response-file detection
# ─────────────────────────────────────────────────────────────────────────────

def resolve_response_file(directory: str, identifier: str, claim_key: str, suffix: str) -> str:
    """
    Locate the response JSON for a claim inside ``directory``.

    Strategy (first hit wins):
      1. exact  <identifier><suffix>            (current karate naming)
      2. exact  <claim_key><suffix>
      3. scan   any *<suffix> file whose name also contains the claim key
      4. scan   any file containing the claim key + a source token from suffix
    Falls back to the canonical path (1) so the caller's existing
    ``os.path.exists`` "response missing" handling still triggers unchanged.
    """
    canonical = os.path.join(directory, f"{identifier}{suffix}")
    if os.path.exists(canonical):
        return canonical

    alt = os.path.join(directory, f"{claim_key}{suffix}")
    if os.path.exists(alt):
        return alt

    if os.path.isdir(directory):
        # source token e.g. "_alexResponse.json" → "alex"
        token = re.sub(r'[^a-z0-9]', '', suffix.lower())
        for fname in os.listdir(directory):
            low = fname.lower()
            if claim_key and claim_key.lower() in low and low.endswith(suffix.lower()):
                return os.path.join(directory, fname)
        for fname in os.listdir(directory):
            low = fname.lower()
            if claim_key and claim_key.lower() in low and token[:4] in low and low.endswith('.json'):
                return os.path.join(directory, fname)

    return canonical


# ─────────────────────────────────────────────────────────────────────────────
# 5. Dynamic report naming
# ─────────────────────────────────────────────────────────────────────────────

def build_report_name(consumer_name: str, config: dict, resolved: dict) -> str:
    """
    Build a descriptive, dynamic report file name from the run context, e.g.
        PPKG_Test_vs_QAE_SummaryAPI_20260807_162307.xlsx

    Convention: ``<Consumer>_<SourceEnv>_vs_<TargetEnv>_<RequestType>_<Timestamp>``
    where the tokens are taken from metadata parsed off the feature files.
    A custom template may still be supplied via ``config['report_name_template']``.
    Falls back to the legacy ``<consumer>_validation_report_<ts>.xlsx``.
    """
    now = datetime.now()

    # Consumer token: prefer the expected-source label (e.g. "PPKG"), else the
    # configured consumer_name's first word, else the registration name.
    consumer_token = (
        resolved.get('expected_source')
        or (config.get('consumer_name') or consumer_name).split()[0]
    )
    request_type = (resolved.get('request_type') or '').replace(' ', '') or 'Validation'

    ctx = {
        'consumer':        consumer_name,
        'consumer_token':  consumer_token,
        'module':          resolved.get('module') or config.get('claim_type', ''),
        'expected':        resolved.get('expected_source', 'PPKG'),
        'actual':          resolved.get('actual_source', 'Alex'),
        'validation_type': resolved.get('validation_type', ''),
        'search_type':     resolved.get('search_type', '') or resolved.get('validation_type', ''),
        'source_env':      resolved.get('source_env', '') or 'SRC',
        'target_env':      resolved.get('target_env', '') or 'TGT',
        'request_type':    request_type,
        'environment':     resolved.get('environment', '') or 'ENV',
        'date':            now.strftime('%Y%m%d'),
        'timestamp':       now.strftime('%Y%m%d_%H%M%S'),
    }

    template = config.get('report_name_template')
    if template:
        try:
            base = template.format(**ctx)
        except (KeyError, IndexError):
            base = f"{consumer_name}_validation_report_{ctx['timestamp']}"
    else:
        parts = [p for p in (
            ctx['consumer_token'], ctx['source_env'], 'vs', ctx['target_env'],
            ctx['request_type'], ctx['timestamp'],
        ) if p]
        base = '_'.join(str(p) for p in parts)

    if not base.lower().endswith('.xlsx'):
        base += '.xlsx'
    # sanitise for filesystem safety
    return re.sub(r'[^\w.\-()]+', '_', base)


# ─────────────────────────────────────────────────────────────────────────────
# 6. Report header metadata + injection into the workbook
# ─────────────────────────────────────────────────────────────────────────────

def build_report_metadata(consumer_name: str, config: dict, resolved: dict) -> List[Tuple[str, str]]:
    """
    Build the ordered (label, value) rows for the report header. Pulls from the
    config/feature metadata resolved during discovery.
    """
    meta = dict(config.get('metadata') or {})
    meta.update({k: v for k, v in resolved.items() if v})

    mapping_names = ', '.join(
        os.path.basename(m) for m in (config.get('mapping_sheets') or [])
    )

    rows = [
        ('Validation Name', meta.get('validation_name') or consumer_name),
        ('Environment',     meta.get('environment', '')),
        ('Feature Name',    meta.get('feature_name', '')),
        ('Scenario Name',   meta.get('scenario_name', '')),
        ('Validation Type', meta.get('validation_type', '')),
        ('API URL',         meta.get('api_url', '')),
        ('Expected Source', meta.get('expected_source', 'PPKG')),
        ('Actual Source',   meta.get('actual_source', 'Alex')),
        ('Mapping Sheet',   mapping_names),
        ('Execution Time',  datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
    ]
    return rows


def inject_report_info_sheet(report_path: str, header_rows: List[Tuple[str, str]]) -> None:
    """
    Add a ``Report Info`` sheet (as the first tab) containing the header
    metadata. This is purely additive — it opens the already-written workbook
    and inserts a new sheet, leaving every existing sheet untouched.
    """
    if not header_rows or not os.path.exists(report_path):
        return
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except Exception:
        return

    try:
        wb = load_workbook(report_path)
    except Exception:
        return

    if 'Report Info' in wb.sheetnames:
        del wb['Report Info']
    ws = wb.create_sheet('Report Info', 0)

    title_fill = PatternFill('solid', fgColor='1F4E79')
    label_fill = PatternFill('solid', fgColor='D9E2F3')
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells('A1:B1')
    tcell = ws.cell(row=1, column=1, value='Validation Report — Header Information')
    tcell.font = Font(bold=True, size=13, color='FFFFFF')
    tcell.fill = title_fill
    tcell.alignment = Alignment(horizontal='left', vertical='center')

    for i, (label, value) in enumerate(header_rows, start=3):
        lcell = ws.cell(row=i, column=1, value=label)
        vcell = ws.cell(row=i, column=2, value=('' if value is None else str(value)))
        lcell.font = Font(bold=True)
        lcell.fill = label_fill
        lcell.border = border
        vcell.border = border
        vcell.alignment = Alignment(vertical='center', wrap_text=False)

    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 90

    try:
        wb.save(report_path)
    except Exception:
        pass

