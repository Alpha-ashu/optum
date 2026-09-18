"""
Shared Validation Framework (project-wide)
===========================================
Common, dependency-free helpers reused by EVERY comparison/validation script in
this project (DMV Index Schema Validation, DMV_Index_Schema_Validation, upm_ppkg_validation,
beyond_compare_validation, …).

This module intentionally has NO import dependency on any single validation
script, so any script can do:

    sys.path.append(os.path.join(os.path.dirname(__file__), '..'))
    from shared_validation_framework import (
        parse_feature_file, build_source_target_metadata,
        classify_missing_field, build_report_name, build_report_metadata,
        inject_report_info_sheet,
    )

Responsibilities
----------------
1. Parse a Karate ``.feature`` file for metadata (Feature title, Scenario,
   API URL, Environment) — feature-file-driven, never hardcoded.
2. Provide ONE shared Blocker/Non-Blocker classification rule for "field
   present on one side, genuinely absent on the other" so every validation
   script in the project applies the identical policy:
       field exists in SOURCE, absent in TARGET -> "Missing In <Target>", Blocker
       field exists in TARGET, absent in SOURCE -> "Missing In <Source>", Blocker
   unless the field is present in an explicit, per-script IGNORABLE list.
3. Build a dynamic, descriptive report file name and a "Report Info" style
   metadata block driven entirely by parsed feature-file metadata.

Nothing here reads or mutates comparison results directly — each script stays
responsible for its own field-by-field diff logic; this module only supplies
the shared *policy* (classification + metadata + naming) so behaviour is
consistent project-wide.
"""

from __future__ import annotations

import os
import re
from datetime import datetime
from typing import Dict, List, Optional, Tuple


# ─────────────────────────────────────────────────────────────────────────────
# 1. Feature-file metadata extraction (Feature -> Runner -> Engine -> Report)
# ─────────────────────────────────────────────────────────────────────────────

# Canonical environment labels for known upstream/environment tokens. Kept in
# sync with DMV_Index_Schema_Validation/validation_framework.py so every
# pipeline reports environment names identically.
ENV_LABELS: Dict[str, str] = {
    'qae': 'QAE', 'uat': 'UAT', 'prod': 'PROD', 'production': 'PROD',
    'stage': 'STAGE', 'staging': 'STAGE', 'test': 'Test', 'tst': 'Test',
    'dev': 'DEV', 'development': 'Development', 'nonprod': 'NONPROD',
    'non-prod': 'NONPROD', 'stagecloud': 'STAGE',
}


def _label_env(value: str) -> str:
    v = (value or '').strip()
    if not v:
        return ''
    return ENV_LABELS.get(v.lower(), v if len(v) <= 4 else v.title())


def _extract_environment(text: str, api_url: str) -> str:
    """Same priority order as the DMV framework: X-Upstream-Env header, then
    an explicit ``def env``/``def environment``, then URL/text heuristics."""
    m = re.search(r"""X-Upstream-Env\s*=\s*['"]([^'"]+)['"]""", text, re.IGNORECASE)
    if m:
        return _label_env(m.group(1))
    m = re.search(r"""def\s+(?:env|environment)\s*=\s*['"]([^'"]+)['"]""", text, re.IGNORECASE)
    if m:
        return _label_env(m.group(1))
    haystack = f"{api_url} {text}".lower()
    for needle, label in (
        ('nonprod', 'NONPROD'), ('non-prod', 'NONPROD'), ('qae', 'QAE'),
        ('stage', 'STAGE'), ('staging', 'STAGE'), ('uat', 'UAT'),
        ('prod', 'PROD'), ('/dev/', 'DEV'),
    ):
        if needle in haystack:
            return label
    return ''


def _extract_request_type(text: str) -> str:
    """Best-effort HTTP verb detection (GET / POST / PUT / DELETE)."""
    m = re.search(r'\bWhen\s+method\s+(\w+)', text, re.IGNORECASE)
    if m:
        return m.group(1).upper()
    for verb in ('post', 'get', 'put', 'delete', 'patch'):
        if re.search(rf'\bmethod\s+{verb}\b', text, re.IGNORECASE):
            return verb.upper()
    return ''


def parse_feature_file(feature_path: Optional[str], base_dir: Optional[str] = None) -> Dict[str, str]:
    """Parse a Karate ``.feature`` file and return best-effort metadata:
    feature_name, scenario_name, api_url, environment, request_type.
    Tolerant of missing files — never raises."""
    meta: Dict[str, str] = {}
    if not feature_path:
        return meta
    if base_dir and not os.path.isabs(feature_path):
        feature_path = os.path.join(base_dir, feature_path)
    if not os.path.exists(feature_path):
        return meta
    try:
        with open(feature_path, 'r', encoding='utf-8', errors='replace') as f:
            text = f.read()
    except OSError:
        return meta

    m = re.search(r'\s*Feature\s*:\s*(.+)', text)
    if m:
        meta['feature_name'] = m.group(1).strip()

    m = re.search(r'\s*Scenario(?:\s+Outline)?\s*:\s*(.+)', text)
    if m:
        meta['scenario_name'] = re.sub(r'\s+', ' ', m.group(1)).strip()

    m = re.search(r"""def\s+api_url\s*=\s*['"]([^'"]+)['"]""", text, re.IGNORECASE)
    if not m:
        m = re.search(r"""Given\s+url\s+['"]([^'"]+)['"]""", text, re.IGNORECASE)
    if not m:
        m = re.search(r"""['"](https?://[^'"]+)['"]""", text)
    if m:
        meta['api_url'] = m.group(1).strip()

    env = _extract_environment(text, meta.get('api_url', ''))
    if env:
        meta['environment'] = env

    req = _extract_request_type(text)
    if req:
        meta['request_type'] = req

    return meta


def build_source_target_metadata(
    source_feature_file: Optional[str],
    target_feature_file: Optional[str],
    base_dir: Optional[str] = None,
) -> Dict[str, str]:
    """Parse both SOURCE (expected/source-of-truth) and TARGET (actual) feature
    files and return a flat metadata dict consumed by the report builders."""
    src = parse_feature_file(source_feature_file, base_dir)
    tgt = parse_feature_file(target_feature_file, base_dir)
    return {
        'source_url': src.get('api_url', ''),
        'target_url': tgt.get('api_url', ''),
        'source_env': src.get('environment', ''),
        'target_env': tgt.get('environment', ''),
        'source_feature_name': src.get('feature_name', ''),
        'target_feature_name': tgt.get('feature_name', ''),
        'request_type': src.get('request_type') or tgt.get('request_type', ''),
    }


# ─────────────────────────────────────────────────────────────────────────────
# 2. Shared missing-field Blocker/Non-Blocker classification policy
# ─────────────────────────────────────────────────────────────────────────────

def is_ignorable_missing_field(ignorable_fields: List[str], *paths: str) -> bool:
    """True if ANY of `paths` matches an entry in `ignorable_fields`
    (case-insensitive substring match). Empty list => nothing is ignored."""
    if not ignorable_fields:
        return False
    haystacks = [str(p or '').lower() for p in paths]
    for needle in ignorable_fields:
        n = str(needle).lower().strip()
        if n and any(n in h for h in haystacks):
            return True
    return False


def classify_missing_field(
    field_path: str,
    present_in_source: bool,
    present_in_target: bool,
    source_label: str,
    target_label: str,
    ignorable_fields: Optional[List[str]] = None,
) -> Tuple[str, str]:
    """
    Single, project-wide policy for classifying a field that is missing on one
    side. Returns (Category, Severity).

    - present_in_source=True, present_in_target=False -> Missing In <target>, Blocker
      (unless field_path is in ignorable_fields -> Non-Blocker)
    - present_in_target=True, present_in_source=False -> Missing In <source>, Blocker
      (unless ignored)
    - both missing -> "No Value", Non-Blocker (nothing to compare)
    """
    ignorable_fields = ignorable_fields or []
    if present_in_source and not present_in_target:
        if is_ignorable_missing_field(ignorable_fields, field_path):
            return f'Missing In {target_label} (Ignored)', 'Non-Blocker'
        return f'Missing In {target_label} (True Missing)', 'Blocker'
    if present_in_target and not present_in_source:
        if is_ignorable_missing_field(ignorable_fields, field_path):
            return f'Missing In {source_label} (Ignored)', 'Non-Blocker'
        return f'Missing In {source_label} (True Missing)', 'Blocker'
    return 'No Value', 'Non-Blocker'


# ─────────────────────────────────────────────────────────────────────────────
# 3. Dynamic report naming + metadata header (shared across all scripts)
# ─────────────────────────────────────────────────────────────────────────────

def build_report_name(consumer_name: str, resolved: Dict[str, str]) -> str:
    """Build e.g. ``PPKG_Test_vs_QAE_SummaryAPI_20260807_180422.xlsx``."""
    now = datetime.now()
    request_type = (resolved.get('request_type') or '').replace(' ', '') or 'Validation'
    parts = [
        p for p in (
            consumer_name,
            resolved.get('source_env') or 'SRC',
            'vs',
            resolved.get('target_env') or 'TGT',
            request_type,
            now.strftime('%Y%m%d_%H%M%S'),
        ) if p
    ]
    base = '_'.join(str(p) for p in parts) + '.xlsx'
    return re.sub(r'[^\w.\-()]+', '_', base)


def build_report_metadata_rows(consumer_name: str, resolved: Dict[str, str]) -> List[Tuple[str, str]]:
    """Ordered (label, value) rows for a 'Report Info' sheet."""
    return [
        ('Consumer Name',      consumer_name),
        ('Source URL',         resolved.get('source_url', '')),
        ('Target URL',         resolved.get('target_url', '')),
        ('Source Environment', resolved.get('source_env', '')),
        ('Target Environment', resolved.get('target_env', '')),
        ('Request Type',       resolved.get('request_type', '')),
        ('Claim Search Type',  resolved.get('search_type', '')),
        ('Source Feature',     resolved.get('source_feature_name', '')),
        ('Target Feature',     resolved.get('target_feature_name', '')),
        ('Execution Date',     datetime.now().strftime('%Y-%m-%d')),
        ('Execution Time',     datetime.now().strftime('%H:%M:%S')),
    ]


def inject_report_info_sheet(report_path: str, header_rows: List[Tuple[str, str]]) -> None:
    """Add/replace a 'Report Info' sheet as the first tab in an existing
    workbook. Purely additive; no-ops if the workbook can't be opened."""
    if not header_rows or not os.path.exists(report_path):
        return
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except Exception:
        return
    try:
        wb = load_workbook(report_path)
    except Exception:
        return

    if 'Report Info' in wb.sheetnames:
        del wb['Report Info']
    ws = wb.create_sheet('Report Info', 0)

    title_fill = PatternFill('solid', fgColor='1F4E79')
    label_fill = PatternFill('solid', fgColor='D9E2F3')
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells('A1:B1')
    tcell = ws.cell(row=1, column=1, value='Validation Report — Header Information')
    tcell.font = Font(bold=True, size=13, color='FFFFFF')
    tcell.fill = title_fill
    tcell.alignment = Alignment(horizontal='left', vertical='center')

    for i, (label, value) in enumerate(header_rows, start=3):
        lcell = ws.cell(row=i, column=1, value=label)
        vcell = ws.cell(row=i, column=2, value=('' if value is None else str(value)))
        lcell.font = Font(bold=True)
        lcell.fill = label_fill
        lcell.border = border
        vcell.border = border

    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 90

    try:
        wb.save(report_path)
    except Exception:
        pass


def build_blocker_summary(rows_with_severity: List[Dict], severity_key: str = 'Severity',
                           category_key: str = 'Category') -> Dict[str, int]:
    """Given a list of dict rows each carrying a Severity/Category, return a
    {category: count} breakdown PLUS 'Total Blockers' — used to build the
    'Blocker Breakdown' sheet consistently across all scripts."""
    breakdown: Dict[str, int] = {}
    total_blockers = 0
    for row in rows_with_severity:
        if row.get(severity_key) == 'Blocker':
            cat = row.get(category_key, 'Uncategorized')
            breakdown[cat] = breakdown.get(cat, 0) + 1
            total_blockers += 1
    breakdown['Total Blockers'] = total_blockers
    return breakdown

